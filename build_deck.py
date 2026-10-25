#!/usr/bin/env python3
"""Build the SafeHaven STL CAC deck (.pptx).
Run: uv run --with python-pptx --with pillow python assets/build_deck.py
Output: demo/SafeHaven-CAC-deck.pptx  (16:9, on-brand)
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIA = os.path.join(ROOT, 'demo', 'diagrams')
SHOT = os.path.join(ROOT, 'demo')
OUT = os.path.join(ROOT, 'demo', 'SafeHaven-CAC-deck.pptx')

BG = RGBColor(0x0E, 0x1A, 0x2B)
ACCENT = RGBColor(0x4E, 0xA8, 0xDE)
TEXT = RGBColor(0xF2, 0xF6, 0xFC)
MUTED = RGBColor(0x9D, 0xB0, 0xCC)
ACTIVE = RGBColor(0x2E, 0xCC, 0x71)

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W = prs.slide_width


def bg(s):
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG


def tb(s, x, y, w, h, text, size, color=TEXT, bold=False, align=PP_ALIGN.LEFT, ls=1.0):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    for i, ln in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = ls
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = 'Arial'
    return box


def bar(s, x=0.0, y=0.0, w=13.333, h=0.16, color=ACCENT):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color; sh.line.fill.background()


def pic(s, path, x, y, w, h):
    iw, ih = Image.open(path).size; ar = iw / ih
    if w / h > ar:
        nh = h; nw = h * ar
    else:
        nw = w; nh = w / ar
    s.shapes.add_picture(path, Inches(x + (w - nw) / 2), Inches(y + (h - nh) / 2), Inches(nw), Inches(nh))


# 1 Title
s = prs.slides.add_slide(BLANK); bg(s); bar(s, y=0, h=0.2)
p = os.path.join(SHOT, 'screen-active-banner.png')
if os.path.exists(p):
    pic(s, p, 9.3, 1.0, 3.3, 5.5)
tb(s, 0.9, 2.0, 8.0, 1.2, 'SafeHaven STL', 58, TEXT, bold=True)
tb(s, 0.9, 3.2, 8.0, 0.8, 'Find a warming or cooling center — fast, offline, when it matters', 23, ACCENT)
tb(s, 0.9, 4.2, 8.0, 1.4,
   'Every St. Louis warming & cooling center in one place.\nNearest first. One tap to directions or a call. Built for a crisis.',
   20, MUTED, ls=1.2)
tb(s, 0.9, 6.7, 9, 0.4, 'Congressional App Challenge  ·  React Native / Expo / TypeScript', 14, MUTED)

# 2 Problem
s = prs.slides.add_slide(BLANK); bg(s)
tb(s, 0.9, 0.6, 11.5, 0.9, 'The problem', 40, TEXT, bold=True)
bar(s, x=0.9, y=1.5, w=2.2, h=0.06)
tb(s, 0.9, 1.9, 11.5, 4.6,
   'When St. Louis hits dangerous cold or heat, the City opens warming and cooling\n'
   'centers to protect people — the City\u2019s own words: those experiencing\n'
   'homelessness, the elderly, and people with disabilities.\n\n'
   'But that information is scattered across separate web pages, PDFs, press\n'
   'releases, and a phone line. The people who need a safe place most are the\n'
   'same people least able to dig through government websites to find one.\n\n'
   'In extreme weather, that gap isn\u2019t an inconvenience. People die of exposure.',
   23, MUTED, ls=1.25)

# 3 Not another directory
s = prs.slides.add_slide(BLANK); bg(s)
tb(s, 0.9, 0.6, 11.5, 0.9, 'Not another directory', 40, TEXT, bold=True)
bar(s, x=0.9, y=1.5, w=2.2, h=0.06)
tb(s, 0.9, 1.9, 11.5, 1.0,
   'Broad services directories exist (findhelp, CompassionSTL). SafeHaven is\n'
   'different — narrow, local, and built for the moment of crisis:',
   22, MUTED, ls=1.2)
tb(s, 0.9, 3.3, 11.5, 3.4,
   '• No single St. Louis app aggregates the warming/cooling centers — SafeHaven does\n\n'
   '• LIVE cooling data from the Missouri DHSS feed + an "open now" filter —\n'
   '   so you don\u2019t set out for a center that\u2019s already closed\n\n'
   '• The others assume a connection; SafeHaven works OFFLINE on a curated fallback\n\n'
   '• Built for the person outside with a dying phone — two giant buttons, no typing',
   20, TEXT, ls=1.3)

# 4 Screens
s = prs.slides.add_slide(BLANK); bg(s)
tb(s, 0.9, 0.5, 11.5, 0.9, 'What it does', 40, TEXT, bold=True)
bar(s, x=0.9, y=1.4, w=2.2, h=0.06)
shots = [('screen-active-banner.png', 'Find + emergency banner'),
         ('screen-nearme.png', 'Near me, nearest first'),
         ('screen-detail.png', 'Directions & call'),
         ('screen-settings.png', 'Theme · offline')]
x = 0.7
for fn, cap in shots:
    pp = os.path.join(SHOT, fn)
    if os.path.exists(pp):
        pic(s, pp, x, 1.7, 2.9, 4.6)
        tb(s, x, 6.35, 2.9, 0.5, cap, 15, ACCENT, bold=True, align=PP_ALIGN.CENTER)
    x += 3.1

# 5-7 diagrams
for fn in ['01-data.png', '02-crisis.png', '03-arch.png']:
    s = prs.slides.add_slide(BLANK); bg(s)
    pp = os.path.join(DIA, fn)
    if os.path.exists(pp):
        pic(s, pp, 0.5, 0.4, 12.333, 6.7)

# 8 Close
s = prs.slides.add_slide(BLANK); bg(s); bar(s, y=0, h=0.2)
tb(s, 0.9, 1.5, 11.5, 1.0, 'Help, in seconds — for the people who need it', 36, TEXT, bold=True)
tb(s, 0.9, 2.8, 11.5, 2.6,
   'A real, installable app: live Missouri DHSS cooling data, an "open now" filter,\n'
   'and a curated offline fallback — working on any phone, in any weather.\n\n'
   'Next: live warming activation (Code Blue), community hour corrections, nearby\n'
   'food + overnight shelters, and multilingual support.',
   22, MUTED, ls=1.3)
tb(s, 0.9, 6.5, 11.5, 0.5,
   'A community reference, not an official City service. Call 2-1-1 to confirm.', 15, ACCENT)

prs.save(OUT)
print('Wrote', OUT, '-', len(prs.slides._sldIdLst), 'slides')
